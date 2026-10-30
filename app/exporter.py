"""Экспорт расшифровки в TXT, SRT, DOCX и сохранение/открытие проекта."""

import json
import os

from app.models import Project, format_time


class ExportError(Exception):
    """Не получилось сохранить или открыть файл. Текст показывается пользователю."""


def _check_segments(segments):
    if not segments:
        raise ExportError("Нечего сохранять: сначала расшифруйте запись.")


def segments_to_text(segments, with_timecodes=True):
    """Собирает сегменты в обычный текст (с таймкодами или без)."""
    if with_timecodes:
        return "\n".join(
            f"[{format_time(s.start)} – {format_time(s.end)}] {s.text}" for s in segments
        )
    return " ".join(s.text for s in segments)


def export_txt(segments, path, with_timecodes=True):
    _check_segments(segments)
    _write(path, segments_to_text(segments, with_timecodes) + "\n")


def segments_to_srt(segments):
    """Формат субтитров SRT: номер, время, текст и пустая строка."""
    blocks = []
    for i, s in enumerate(segments, start=1):
        blocks.append(f"{i}\n{format_time(s.start, srt=True)} --> {format_time(s.end, srt=True)}\n{s.text}\n")
    return "\n".join(blocks)


def export_srt(segments, path):
    _check_segments(segments)
    _write(path, segments_to_srt(segments))


def export_docx(segments, path, title="Расшифровка аудиозаписи", source_name="",
                with_timecodes=True):
    _check_segments(segments)
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except ImportError as exc:
        raise ExportError("Для экспорта в Word установите библиотеку python-docx.") from exc

    doc = Document()
    doc.add_heading(title, level=1)
    if source_name:
        doc.add_paragraph(f"Исходный файл: {source_name}")
    for s in segments:
        p = doc.add_paragraph()
        if with_timecodes:
            run = p.add_run(f"[{format_time(s.start)} – {format_time(s.end)}]  ")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        p.add_run(s.text)
    try:
        doc.save(path)
    except PermissionError as exc:
        raise ExportError("Нет доступа к файлу. Возможно, он открыт в Word — закройте его.") from exc
    except OSError as exc:
        raise ExportError(f"Не удалось сохранить файл:\n{exc}") from exc


def save_project(project, path):
    _check_segments(project.segments)
    _write(path, json.dumps(project.to_dict(), ensure_ascii=False, indent=2))


def load_project(path):
    if not os.path.isfile(path):
        raise ExportError(f"Файл проекта не найден:\n{path}")
    try:
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        return Project.from_dict(data)
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise ExportError("Файл проекта повреждён или имеет неверный формат.") from exc
    except (ValueError, KeyError, TypeError) as exc:
        raise ExportError(f"Файл проекта содержит ошибку: {exc}") from exc


def _write(path, text):
    try:
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)
    except PermissionError as exc:
        raise ExportError("Нет доступа к файлу. Возможно, он открыт в другой программе.") from exc
    except OSError as exc:
        raise ExportError(f"Не удалось сохранить файл:\n{exc}") from exc
